#!/usr/bin/env python3
"""backends/pptx.py - FigSmith IR -> PPTX slides (python-pptx, no browser).

One IR, one slide. The geometry comes from the same resolver the SVG backend uses, so
an SVG figure and its PPTX twin place every box at the same coordinates.

Usage:
    python -m backends.pptx <ir.json> -o out.pptx [--slide-w 13.333] [--slide-h 7.5]
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backends.svg import resolve_layout  # noqa: E402
from core.tokens import load as load_tokens  # noqa: E402

EMU_PER_INCH = 914400
# PPTX colors are RGB hex without '#'


def _rgb(hex_color: str):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _render_statistical(ir: dict, out_path: str | Path, slide_w_in: float,
                        slide_h_in: float, base_dir: Path | None = None) -> str:
    """A statistical IR has no topology, so the structure path would ship a blank slide.

    The chart is rendered by matplotlib (the only sanctioned route for data) and placed
    as the slide's figure, with the title and axis labels as real, editable text. That
    keeps the four-way promise honest: the PPTX carries the same content as the SVG.
    """
    import tempfile

    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    from backends.mpl2svg import render as render_mpl

    tokens = load_tokens(ir.get("meta", {}).get("tokens", "academic"))
    title = ir.get("meta", {}).get("title") or ir.get("content", {}).get("title") or ""
    axis = ir.get("content", {}).get("axis", {})
    legend = [s.get("name") for s in ir.get("data", {}).get("series", []) if s.get("name")]

    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    margin_in = 0.5
    top = margin_in
    if title:
        box = slide.shapes.add_textbox(Inches(margin_in), Inches(margin_in),
                                       Inches(slide_w_in - 2 * margin_in), Inches(0.7))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = str(title)
        run.font.name = tokens.font("figure_title", "latin")
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = _rgb(tokens.c("colors.ink"))
        top += 0.9

    with tempfile.TemporaryDirectory() as tmp:
        chart_png = Path(tmp) / "chart.png"
        # base_dir: resolve data.source relative to the IR document, not the cwd.
        render_mpl(ir, out_png=str(chart_png), dpi=300, base_dir=base_dir)
        from PIL import Image

        w_px, h_px = Image.open(chart_png).size
        avail_w = slide_w_in - 2 * margin_in
        avail_h = slide_h_in - top - margin_in
        fit = min(avail_w / w_px, avail_h / h_px)
        pic_w = w_px * fit
        pic_h = h_px * fit
        left = (slide_w_in - pic_w) / 2
        slide.shapes.add_picture(str(chart_png), Inches(left), Inches(top),
                                 Inches(pic_w), Inches(pic_h))

    # Axis labels and series names are stated as text so the slide still reads correctly
    # when the image is missing, and so a text-level content check can see them. Each
    # item is its own run, so a checker sees the label itself rather than "x: label".
    captions = [str(v) for _, v in axis.items() if str(v).strip()]
    captions += [str(s.get("name")) for s in ir.get("data", {}).get("series", [])
                 if s.get("name")]
    if captions:
        cap = slide.shapes.add_textbox(Inches(margin_in), Inches(slide_h_in - margin_in - 0.5),
                                       Inches(slide_w_in - 2 * margin_in), Inches(0.5))
        tf = cap.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        for i, caption in enumerate(captions):
            if i:
                sep = p.add_run()
                sep.text = "  |  "
                sep.font.name = tokens.font("annotation", "latin")
                sep.font.size = Pt(tokens.font_size("annotation"))
                sep.font.color.rgb = _rgb(tokens.c("colors.ink_tertiary"))
            run = p.add_run()
            run.text = caption
            run.font.name = tokens.font("annotation", "latin")
            run.font.size = Pt(tokens.font_size("annotation"))
            run.font.color.rgb = _rgb(tokens.c("colors.ink_secondary"))

    prs.save(str(out_path))
    return str(out_path)


def render(ir: dict, out_path: str | Path, slide_w_in: float = 13.333,
           slide_h_in: float = 7.5, base_dir: Path | None = None) -> str:
    if ir.get("data", {}).get("series"):
        return _render_statistical(ir, out_path, slide_w_in, slide_h_in, base_dir)

    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    tokens = load_tokens(ir.get("meta", {}).get("tokens", "academic"))
    scale = float(ir.get("layout", {}).get("canvas", {}).get("px_per_unit", 12.0))
    nodes, info = resolve_layout(ir, scale)
    labels = ir.get("content", {}).get("labels", {})
    emphasis = set(ir.get("style", {}).get("emphasis", []))

    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Scale the figure to fit the slide, centered, preserving aspect.
    fig_w_px, fig_h_px = info["canvas_w"], info["canvas_h"]
    margin_in = 0.5
    avail_w = slide_w_in - 2 * margin_in
    avail_h = slide_h_in - 2 * margin_in
    fit = min(avail_w / fig_w_px, avail_h / fig_h_px)
    origin_x = (slide_w_in - fig_w_px * fit) / 2
    origin_y = (slide_h_in - fig_h_px * fit) / 2

    def to_emu(px: float) -> int:
        return int(px * fit * EMU_PER_INCH)

    # bands behind boxes, matching the SVG painter's order
    ordered = [n for n in nodes if n.kind == "band"] + [n for n in nodes if n.kind != "band"]
    for node in ordered:
        if node.kind == "label":
            continue
        shape = slide.shapes.add_shape(
            1,  # MSO_SHAPE.RECTANGLE == 1; kept literal to avoid the enum import
            Emu(to_emu(node.x) + int(origin_x * EMU_PER_INCH)),
            Emu(to_emu(node.y) + int(origin_y * EMU_PER_INCH)),
            Emu(to_emu(node.w)), Emu(to_emu(node.h)),
        )
        if node.kind == "band":
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(tokens.c("colors.surface_band"))
            shape.line.fill.background()
            shape.shadow.inherit = False
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(tokens.c("colors.surface_band"))
            stroke_hex = tokens.c("colors.border")
            if node.id in emphasis:
                accent = tokens.get("colors.accent")
                stroke_hex = accent[0] if isinstance(accent, list) and accent else stroke_hex
            shape.line.color.rgb = _rgb(stroke_hex)
            shape.line.width = Pt(1.4 if node.id in emphasis else 0.75)
            shape.shadow.inherit = False

        text = labels.get(node.id, "")
        if text and node.kind != "band":
            tf = shape.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Emu(int(0.06 * EMU_PER_INCH))
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = text
            font = run.font
            font.name = tokens.font("panel_label", "latin")
            font.size = Pt(tokens.font_size("panel_label"))
            font.color.rgb = _rgb(tokens.c("colors.ink"))
        elif text and node.kind == "band":
            tf = shape.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = Emu(int(0.1 * EMU_PER_INCH))
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = text
            run.font.name = tokens.font("panel_label", "latin")
            run.font.size = Pt(tokens.font_size("panel_label"))
            run.font.bold = True
            run.font.color.rgb = _rgb(tokens.c("colors.ink_secondary"))

    # arrows as straight connectors, with their labels as small text boxes
    by_id = {n.id: n for n in nodes}
    for i, edge in enumerate(ir.get("topology", {}).get("edges", [])):
        src, dst = by_id.get(edge["from"]), by_id.get(edge["to"])
        if src is None or dst is None:
            continue
        x1 = Emu(to_emu(src.right if dst.x > src.right else src.x) + int(origin_x * EMU_PER_INCH))
        y1 = Emu(to_emu(src.cy) + int(origin_y * EMU_PER_INCH))
        x2 = Emu(to_emu(dst.x if dst.x > src.right else src.right) + int(origin_x * EMU_PER_INCH))
        y2 = Emu(to_emu(dst.cy) + int(origin_y * EMU_PER_INCH))
        if dst.y > src.bottom or src.y > dst.bottom:  # vertical relationship
            x1 = Emu(to_emu(src.cx) + int(origin_x * EMU_PER_INCH))
            x2 = Emu(to_emu(dst.cx) + int(origin_x * EMU_PER_INCH))
            y1 = Emu(to_emu(src.bottom if dst.y > src.bottom else src.y) + int(origin_y * EMU_PER_INCH))
            y2 = Emu(to_emu(dst.y if dst.y > src.bottom else dst.bottom) + int(origin_y * EMU_PER_INCH))
        conn = slide.shapes.add_connector(1, x1, y1, x2, y2)  # straight
        conn.line.color.rgb = _rgb(tokens.c("colors.border"))
        conn.line.width = Pt(1.0)

        label_text = labels.get(edge.get("label", ""))
        if label_text:
            # The label rides the midpoint. A tiny borderless text box is the only
            # portable way to carry edge text into PPTX; connectors cannot hold text.
            mx = (int(x1) + int(x2)) // 2
            my = (int(y1) + int(y2)) // 2
            box = slide.shapes.add_textbox(Emu(mx - int(0.5 * EMU_PER_INCH)),
                                           Emu(my - int(0.15 * EMU_PER_INCH)),
                                           Emu(int(1.0 * EMU_PER_INCH)),
                                           Emu(int(0.3 * EMU_PER_INCH)))
            tf = box.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = label_text
            run.font.name = tokens.font("annotation", "latin")
            run.font.size = Pt(tokens.font_size("annotation"))
            run.font.color.rgb = _rgb(tokens.c("colors.ink_secondary"))

    prs.save(str(out_path))
    return str(out_path)


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description="Render a FigSmith IR to a PPTX slide.")
    ap.add_argument("ir")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--slide-w", type=float, default=13.333)
    ap.add_argument("--slide-h", type=float, default=7.5)
    args = ap.parse_args(argv)

    ir_path = Path(args.ir)
    ir = json.loads(ir_path.read_text(encoding="utf-8-sig"))
    out = args.out or str(ir_path.with_suffix(".pptx"))
    # data.source is resolved relative to the IR so a statistical slide is reproducible.
    path = render(ir, out, args.slide_w, args.slide_h,
                  base_dir=ir_path.resolve().parent)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
